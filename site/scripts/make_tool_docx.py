#!/usr/bin/env python3
"""由 src/data/tools/*.md 用 pandoc 生成 public/templates/*.docx，并补表格边框与中文字体。改 md 后重新运行。"""
import subprocess, pathlib
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt
root = pathlib.Path(__file__).resolve().parent.parent
FONT = '微软雅黑'
for md in sorted((root / 'src/data/tools').glob('*.md')):
    out = root / 'public/templates' / (md.stem + '.docx')
    subprocess.run(['pandoc', str(md), '-f', 'gfm', '-o', str(out)], check=True)
    d = Document(out)
    from docx.oxml import OxmlElement
    for t in d.tables:
        tblPr = t._tbl.tblPr
        for old in tblPr.findall(qn('w:tblBorders')): tblPr.remove(old)
        b = OxmlElement('w:tblBorders')
        for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
            e = OxmlElement('w:' + side); e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), '4'); e.set(qn('w:color'), '888888'); b.append(e)
        tblPr.append(b)
    for st in d.styles:
        try:
            rpr = st.element.get_or_add_rPr(); rpr.rFonts.set(qn('w:eastAsia'), FONT); rpr.rFonts.set(qn('w:ascii'), FONT); rpr.rFonts.set(qn('w:hAnsi'), FONT)
        except Exception: pass
    def fix(p):
        for r in p.runs:
            r._element.get_or_add_rPr(); r.font.name = FONT; r._element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
    for p in d.paragraphs: fix(p)
    for t in d.tables:
        for row in t.rows:
            for c in row.cells:
                for p in c.paragraphs: fix(p)
    d.save(out); print('wrote', out)
