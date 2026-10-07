#!/usr/bin/env python3
"""由 src/data/judgment-card-template.md 生成 public/templates/judgment-card-template.docx。
空白模板在 Word 里做成可填表格，不带 Markdown 符号。改了 md 后重新运行本脚本。"""
import re, pathlib
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.oxml.ns import qn
from docx.enum.text import WD_BREAK

root = pathlib.Path(__file__).resolve().parent.parent
md = (root / 'src/data/judgment-card-template.md').read_text(encoding='utf-8')
FONT = '微软雅黑'

doc = Document()
sec = doc.sections[0]
sec.left_margin = sec.right_margin = Cm(2.2)
sec.top_margin = sec.bottom_margin = Cm(2)
st = doc.styles['Normal']
st.font.name = FONT; st.font.size = Pt(10.5)
st.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)

def run_fmt(r, bold=False, size=None, color=None):
    r.font.name = FONT; r._element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
    r.bold = bold
    if size: r.font.size = Pt(size)
    if color: r.font.color.rgb = RGBColor(*color)

def para(text='', bold=False, size=None, color=None, after=4):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(after)
    # 支持 **粗体**
    for i, seg in enumerate(re.split(r'\*\*(.+?)\*\*', text)):
        if seg: run_fmt(p.add_run(seg), bold=bold or i % 2 == 1, size=size, color=color)
    return p

def h(text):
    p = para(text, bold=True, size=14, color=(0x20,0x1e,0x1d), after=6)
    p.paragraph_format.space_before = Pt(14)

def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    from docx.oxml import OxmlElement
    s = OxmlElement('w:shd'); s.set(qn('w:val'), 'clear'); s.set(qn('w:color'), 'auto'); s.set(qn('w:fill'), hexcolor); tcPr.append(s)

def table(rows, widths, header=None, fill_col=None, min_h=None):
    t = doc.add_table(rows=0, cols=len(widths)); t.style = 'Table Grid'
    if header:
        cells = t.add_row().cells
        for c, x in zip(cells, header):
            c.text = ''; run_fmt(c.paragraphs[0].add_run(x), bold=True); shade(c, 'EAE9E9')
    for row in rows:
        r = t.add_row()
        if min_h: r.height = Cm(min_h)
        for i, (c, x) in enumerate(zip(r.cells, row)):
            c.text = ''
            run_fmt(c.paragraphs[0].add_run(x), bold=(i == 0 and fill_col is not None),
                    color=(0x80,0x80,0x80) if i == 2 and fill_col is not None else None, size=9.5 if i == 2 and fill_col is not None else None)
    for row in t.rows:
        for c, w in zip(row.cells, widths): c.width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

# ---- 标题与用法
para('判断卡 · 标准模板（一页纸）', bold=True, size=20, after=8)
para('用途：给一个**重大、有后果、会被追问**的判断做记录。')
para('用法：复制下面第一节的表格，填完就是一张卡。**填不满，说明这件事还没想清楚，不要下注。**')
para('来源：从判断系统 schema 2.1 的字段里提取核心十二项，去掉内部管理字段。　tommywang.cn/method/judgment-card-template/', color=(0x80,0x80,0x80), size=9)

# ---- 一、空白模板：由 md 的空白模板代码块解析成"字段 | 填写区 | 说明"
h('一、空白模板')
block = re.search(r'```markdown\n(.*?)\n```', md, re.S).group(1)
title_line = block.splitlines()[0].lstrip('# ').strip()
rows = [('卡片编号与标题', '', title_line)]
for sect in re.split(r'\n(?=## )', block.split('\n', 1)[1].strip()):
    lines = sect.strip().splitlines()
    name = lines[0].lstrip('# ').strip()
    hint = '\n'.join(l for l in lines[1:] if l.strip())
    rows.append((name, '', hint))
table(rows, [3.6, 6.4, 6.6], header=['字段', '填写区', '怎么填'], fill_col=0, min_h=1.3)

# ---- 二、自查
h('二、填完自查（8 条，有一条不过就别发出去）')
sec2 = re.search(r'## 二、.*?\n(.*?)\n---', md, re.S).group(1)
chk = [[c.strip().strip('*') for c in l.strip().strip('|').split('|')] for l in sec2.splitlines() if re.match(r'\|\s*\d', l)]
table(chk, [1, 6, 9.6], header=['#', '自查项', '为什么'])

# ---- 三、示例
h('三、填写示例（公开卡）')
ex = re.search(r'```markdown\n(# JUG-TW.*?)\n```', md, re.S).group(1)
for l in ex.splitlines():
    l = l.strip()
    if not l: continue
    if l.startswith('# '): para(l[2:], bold=True, size=11, after=3)
    elif l.startswith('## '): para(l[3:], bold=True, after=1)
    else: para(l, after=1)
para('这张卡是公开的：https://tommywang.cn/research/ledger/', color=(0x80,0x80,0x80), size=9)

# ---- 四、纪律
h('四、四条填写纪律（比模板更重要）')
sec4 = re.search(r'## 四、.*?\n(.*?)\n---', md, re.S).group(1)
for l in sec4.splitlines():
    if re.match(r'\d\.', l.strip()): para(l.strip())

# ---- 五
h('五、不要给每件事建卡')
sec5 = md.split('## 五、不要给每件事建卡')[1]
for l in sec5.splitlines():
    if l.strip(): para(l.strip())

out = root / 'public/templates/judgment-card-template.docx'
doc.save(out); print('wrote', out)
